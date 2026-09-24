# -*- coding: utf-8 -*-
"""Xuất Excel (.xlsx) — tiêu đề + bảng theo cấu hình đối tượng động."""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _side():
    return Side(style="thin")


BORDER = Border(left=_side(), right=_side(), top=_side(), bottom=_side())
BORDER_DOT = Border(
    left=Side(style="dotted"),
    right=Side(style="dotted"),
    top=Side(style="dotted"),
    bottom=Side(style="dotted"),
)


def export_report_xlsx(report: dict, period_label: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Báo cáo"

    dept = report["dept"]
    structure = report["structure"]
    values = report["values"]
    all_cols = sorted(structure["columns"], key=lambda c: (c.sort_order, c.id))
    n_data_cols = max(len(all_cols), 1)

    # ---- Tiêu đề ----
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    ws.cell(1, 1, dept.hospital).font = Font(bold=True, size=12)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=2)
    ws.cell(2, 1, "Số ....../2026")

    title = report["template"].title if report["template"] else "BÁO CÁO"
    ws.merge_cells(start_row=1, start_column=3, end_row=1, end_column=3 + n_data_cols)
    c = ws.cell(1, 3, title)
    c.font = Font(bold=True, size=14)
    c.alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=2, start_column=3, end_row=2, end_column=3 + n_data_cols)
    c = ws.cell(2, 3, period_label)
    c.font = Font(bold=True, size=11)
    c.alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=2)
    ws.cell(3, 1, dept.report_code)
    ws.merge_cells(start_row=3, start_column=3, end_row=3, end_column=3 + n_data_cols)
    c = ws.cell(3, 3, f"Khoa: {dept.name}")
    c.alignment = Alignment(horizontal="center")

    # ---- Header 2 dòng ----
    header_row = 5
    ws.merge_cells(start_row=header_row, start_column=1, end_row=header_row + 1, end_column=1)
    ws.merge_cells(start_row=header_row, start_column=2, end_row=header_row + 1, end_column=2)
    ws.cell(header_row, 1, "STT").alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(header_row, 2, "Danh mục").alignment = Alignment(horizontal="center", vertical="center")

    # gộp nhóm
    seq: list[tuple[str | None, list]] = []
    i = 0
    while i < len(all_cols):
        g = all_cols[i].group_label
        if g:
            j = i
            grp = []
            while j < len(all_cols) and all_cols[j].group_label == g:
                grp.append(all_cols[j])
                j += 1
            seq.append((g, grp))
            i = j
        else:
            seq.append((None, [all_cols[i]]))
            i += 1

    col_idx = 3
    for label, grp_cols in seq:
        if label:
            ws.merge_cells(
                start_row=header_row,
                start_column=col_idx,
                end_row=header_row,
                end_column=col_idx + len(grp_cols) - 1,
            )
            ws.cell(header_row, col_idx, label).alignment = Alignment(horizontal="center")
            for gc in grp_cols:
                ws.cell(header_row + 1, col_idx, gc.label).alignment = Alignment(
                    horizontal="center", wrap_text=True
                )
                col_idx += 1
        else:
            ws.merge_cells(
                start_row=header_row,
                start_column=col_idx,
                end_row=header_row + 1,
                end_column=col_idx,
            )
            ws.cell(header_row, col_idx, grp_cols[0].label).alignment = Alignment(
                horizontal="center", wrap_text=True
            )
            col_idx += 1

    total_cols = col_idx - 1
    for r in range(header_row, header_row + 2):
        for cidx in range(1, total_cols + 1):
            cell = ws.cell(r, cidx)
            cell.border = BORDER
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEBF7")

    # ---- Thân ----
    row_ptr = header_row + 2
    stt = 0
    col_keys_order = [c.col_key for c in all_cols]

    def write_data_row(group_label: str, row_label: str, row_id: int, stt_val=""):
        nonlocal row_ptr
        ws.cell(row_ptr, 1, stt_val).alignment = Alignment(horizontal="center")
        label = f"{group_label} — {row_label}" if group_label else row_label
        ws.cell(row_ptr, 2, label)
        for k, ck in enumerate(col_keys_order):
            v = values.get((row_id, ck), 0)
            cell = ws.cell(row_ptr, 3 + k, int(v) if float(v).is_integer() else v)
            cell.alignment = Alignment(horizontal="center")
            cell.border = BORDER_DOT
        ws.cell(row_ptr, 1).border = BORDER_DOT
        ws.cell(row_ptr, 2).border = BORDER_DOT
        row_ptr += 1

    for sec in structure["sections"]:
        ws.merge_cells(start_row=row_ptr, start_column=1, end_row=row_ptr, end_column=total_cols)
        cell = ws.cell(row_ptr, 1, sec["title"])
        cell.font = Font(bold=True, size=11)
        for cidx in range(1, total_cols + 1):
            ws.cell(row_ptr, cidx).border = BORDER
            ws.cell(row_ptr, cidx).fill = PatternFill("solid", fgColor="F2F2F2")
        row_ptr += 1

        for b in sec["blocks"]:
            ws.merge_cells(start_row=row_ptr, start_column=1, end_row=row_ptr, end_column=total_cols)
            cell = ws.cell(row_ptr, 1, b["label"])
            cell.font = Font(bold=True, italic=True)
            row_ptr += 1
            for r in b["rows"]:
                stt += 1
                write_data_row(r.group_label, r.row_label, r.id, "")
        for r in sec["loose_rows"]:
            stt += 1
            write_data_row(r.group_label, r.row_label, r.id, stt)

    row_ptr += 2
    today = date.today()
    ws.merge_cells(start_row=row_ptr, start_column=1, end_row=row_ptr, end_column=total_cols)
    c = ws.cell(row_ptr, 1, f"Ngày {today.day} tháng {today.month} năm {today.year}")
    c.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 34
    for i in range(3, total_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 9
    ws.print_title_rows = f"{header_row}:{header_row + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
